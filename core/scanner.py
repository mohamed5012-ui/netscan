#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
محرّك مسح الشبكة — netscan core (نسخة 2)
مشترك بين: الويندوز (exe) / الأندرويد (apk) / سطر الأوامر.
مكتبات بايثون القياسية فقط — بدون أي تثبيت.

المزايا:
  • اكتشاف نطاق الشبكة الحقيقي (netmask) والراوتر تلقائيًا
  • دعم الشبكات الكبيرة حتى /16 (65 ألف عنوان)
  • اكتشاف الأجهزة عبر ARP + فحص TCP (يلاقي الأجهزة اللي بتحجب ping)
  • أسماء الأجهزة من 3 مصادر: DNS العكسي + NetBIOS + mDNS
  • قاعدة بيانات الشركات الكاملة (+59 ألف شركة)
  • كشف المشاكل على شبكتك:
      - IP متكرر   : نفس الـ IP رد بأكتر من MAC (تعارض عناوين)
      - MAC مكرر   : نفس الـ MAC ظاهر على أكتر من IP
      - MAC متشابه : ماكات متطابقة ماعدا آخر خانة (غالبًا نفس الجهاز بأكتر من كارت)
      - MAC عشوائي : عنوان خاص/عشوائي (خاصية الخصوصية في الموبايلات الحديثة)
"""

import os
import sys
import re
import gzip
import time
import errno
import queue
import socket
import struct
import ipaddress
import threading
import subprocess

IS_WINDOWS = sys.platform.startswith("win")
IS_MAC = sys.platform == "darwin"
_NO_WINDOW = 0x08000000 if IS_WINDOWS else 0

# أقصى حجم شبكة مسموح (عدد العناوين) — /16
MAX_ADDRESSES = 65536

# منافذ شائعة نجرّب الاتصال بيها لإثبات إن الجهاز حيّ
PROBE_PORTS = [80, 443, 445, 139, 22, 53, 8080, 7547, 23, 8081, 554, 3389]


# ---------------------------------------------------------------------------
#  أدوات أساسية
# ---------------------------------------------------------------------------

def resource_path(rel):
    """مسار ملف مرفق داخل الـ exe (PyInstaller) أو بجانب الكود."""
    for base in (getattr(sys, "_MEIPASS", None),
                 os.path.dirname(os.path.abspath(__file__)),
                 os.path.dirname(os.path.abspath(sys.argv[0])) if sys.argv else "."):
        if not base:
            continue
        p = os.path.join(base, rel)
        if os.path.exists(p):
            return p
    return rel


def normalize_mac(mac):
    mac = mac.lower().replace("-", ":").strip()
    parts = mac.split(":")
    if len(parts) != 6:
        return mac
    return ":".join(p.zfill(2) for p in parts)


def get_local_ip():
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
    except Exception:
        ip = "127.0.0.1"
    finally:
        s.close()
    return ip


def detect_network(prefix_override=None):
    """
    يكتشف نطاق الشبكة الحقيقي للجهاز.
    يحاول يجيب الـ netmask الفعلي للكارت، ولو مقدرش يفترض /24.
    يرجّع (local_ip, network).
    """
    local_ip = get_local_ip()
    if local_ip == "127.0.0.1":
        return local_ip, None

    if prefix_override:
        return local_ip, ipaddress.ip_network(
            f"{local_ip}/{prefix_override}", strict=False)

    # محاولة 1: psutil لو متوفرة
    try:
        import psutil  # type: ignore
        for _n, addrs in psutil.net_if_addrs().items():
            for a in addrs:
                if a.family == socket.AF_INET and a.address == local_ip:
                    nm = a.netmask or "255.255.255.0"
                    return local_ip, ipaddress.ip_network(
                        f"{local_ip}/{nm}", strict=False)
    except Exception:
        pass

    # محاولة 2: قراءة الـ netmask من أوامر النظام
    nm = _netmask_from_system(local_ip)
    if nm:
        try:
            return local_ip, ipaddress.ip_network(
                f"{local_ip}/{nm}", strict=False)
        except Exception:
            pass

    # الافتراضي: /24
    return local_ip, ipaddress.ip_network(f"{local_ip}/24", strict=False)


def _netmask_from_system(local_ip):
    """يحاول استخراج الـ netmask من ifconfig / ip / ipconfig."""
    try:
        if IS_WINDOWS:
            out = subprocess.run(["ipconfig"], capture_output=True, text=True,
                                 creationflags=_NO_WINDOW).stdout
            # نبحث عن السطر اللي فيه الـ IP بتاعنا وبعده Subnet Mask
            lines = out.splitlines()
            for i, line in enumerate(lines):
                if local_ip in line:
                    for j in range(i, min(i + 4, len(lines))):
                        m = re.search(r"(\d+\.\d+\.\d+\.\d+)", lines[j])
                        if m and "255" in lines[j]:
                            return m.group(1)
        else:
            for cmd in (["ip", "-o", "-f", "inet", "addr", "show"],
                        ["ifconfig"]):
                try:
                    out = subprocess.run(cmd, capture_output=True,
                                         text=True).stdout
                except Exception:
                    continue
                # ip: يطلّع صيغة CIDR مباشرة  192.168.1.5/24
                m = re.search(re.escape(local_ip) + r"/(\d+)", out)
                if m:
                    return m.group(1)
    except Exception:
        pass
    return None


# ---------------------------------------------------------------------------
#  قاعدة بيانات الشركات المصنّعة (OUI)
# ---------------------------------------------------------------------------

_OUI_CACHE = None


def load_oui():
    """
    يحمّل قاعدة الشركات من vendors.gz المرفقة (مضغوطة، +59 ألف).
    لو مش موجودة يستخدم الخريطة المصغّرة الاحتياطية.
    يرجّع dict {prefix_hex: vendor}.
    """
    global _OUI_CACHE
    if _OUI_CACHE is not None:
        return _OUI_CACHE

    oui = {}
    try:
        path = resource_path("vendors.gz")
        with gzip.open(path, "rt", encoding="utf-8", errors="ignore") as f:
            for line in f:
                line = line.rstrip("\n")
                if not line or "\t" not in line:
                    continue
                key, name = line.split("\t", 1)
                oui[key.strip().upper()] = name.strip()
    except Exception:
        pass

    if not oui:
        oui = FALLBACK_OUI.copy()
    _OUI_CACHE = oui
    return oui


def vendor_of(mac, oui=None):
    """
    يبحث عن الشركة بالمطابقة الأطول: /36 ثم /28 ثم /24.
    """
    if not mac or len(mac) < 8:
        return ""
    if oui is None:
        oui = load_oui()
    hexs = mac.upper().replace(":", "")
    for nibbles in (9, 7, 6):  # /36، /28، /24
        if len(hexs) >= nibbles:
            v = oui.get(hexs[:nibbles])
            if v:
                return v
    return ""


# خريطة احتياطية مصغّرة (لو ملف vendors.gz مش موجود)
FALLBACK_OUI = {
    "3C5AB4": "Google", "001CB3": "Apple", "A45E60": "Apple", "F01898": "Apple",
    "DCA632": "Raspberry Pi", "B827EB": "Raspberry Pi",
    "000C42": "MikroTik", "4C5E0C": "MikroTik", "488F5A": "MikroTik",
    "2CC81B": "MikroTik", "CC2DE0": "MikroTik", "744D28": "MikroTik",
    "E091F5": "Samsung", "3423BA": "Samsung", "88366C": "Samsung",
    "009ACD": "Huawei", "487B6B": "Huawei", "286ED4": "Huawei",
    "F48CEB": "Xiaomi", "640980": "Xiaomi", "50EC50": "Xiaomi",
    "50C7BF": "TP-Link", "A42BB0": "TP-Link", "E8DE27": "TP-Link",
    "2C56DC": "ASUSTek", "001E8C": "ASUSTek",
}


# ---------------------------------------------------------------------------
#  جدول ARP
# ---------------------------------------------------------------------------

def read_arp_table():
    """يقرأ جدول ARP من النظام ويرجّع list من (ip, mac) — يحتفظ بالتكرارات."""
    pairs = []
    text = ""
    for cmd in (["arp", "-a"], ["ip", "neigh"]):
        try:
            res = subprocess.run(cmd, capture_output=True, text=True,
                                 creationflags=_NO_WINDOW)
            text = res.stdout or ""
            if text.strip():
                break
        except Exception:
            continue

    if not text.strip():
        # أندرويد / لينكس: الملف المباشر
        try:
            with open("/proc/net/arp", "r") as f:
                text = f.read()
        except Exception:
            pass

    mac_re = re.compile(r"(([0-9a-fA-F]{1,2}[:-]){5}[0-9a-fA-F]{1,2})")
    ip_re = re.compile(r"(\d{1,3}\.\d{1,3}\.\d{1,3}\.\d{1,3})")
    for line in text.splitlines():
        mm = mac_re.search(line)
        im = ip_re.search(line)
        if mm and im:
            mac = normalize_mac(mm.group(1))
            if mac in ("00:00:00:00:00:00", "ff:ff:ff:ff:ff:ff"):
                continue
            pairs.append((im.group(1), mac))
    return pairs


# ---------------------------------------------------------------------------
#  إثارة الأجهزة وفحصها
# ---------------------------------------------------------------------------

def stir(ip):
    """يثير الجهاز (UDP + TCP) عشان يظهر في جدول ARP ونعرف إنه حيّ."""
    ip = str(ip)
    alive = False
    try:
        u = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        u.settimeout(0.12)
        u.sendto(b"\x00", (ip, 40000))
        u.close()
    except Exception:
        pass
    for port in PROBE_PORTS[:6]:
        try:
            t = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
            t.settimeout(0.18)
            rc = t.connect_ex((ip, port))
            t.close()
            if rc == 0 or rc == errno.ECONNREFUSED:
                alive = True
                break
        except Exception:
            continue
    return alive


# ---------------------------------------------------------------------------
#  أسماء الأجهزة: DNS عكسي + NetBIOS + mDNS
# ---------------------------------------------------------------------------

def name_reverse_dns(ip):
    try:
        return socket.gethostbyaddr(ip)[0]
    except Exception:
        return ""


def name_netbios(ip, timeout=0.4):
    """
    استعلام NetBIOS Node Status (UDP 137) — بيجيب اسم أجهزة ويندوز.
    """
    try:
        # NBSTAT query لاسم '*'
        tid = b"\xA2\x48"
        header = tid + b"\x00\x00\x00\x01\x00\x00\x00\x00\x00\x00"
        # الاسم المشفّر لـ '*' (NBSTAT)
        enc = b"\x20\x43\x4b" + b"\x41" * 30 + b"\x00"
        question = enc + b"\x00\x21\x00\x01"
        pkt = header + question

        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.settimeout(timeout)
        s.sendto(pkt, (str(ip), 137))
        data, _ = s.recvfrom(1024)
        s.close()

        # عدد الأسماء بعد الهيدر (12) + إعادة السؤال
        off = 12
        # نتخطّى حقل الاسم في الرد
        while off < len(data) and data[off] != 0:
            off += data[off] + 1
        off += 1 + 2 + 2 + 4 + 2  # null + type + class + ttl + rdlength
        if off >= len(data):
            return ""
        num = data[off]
        off += 1
        for _ in range(num):
            if off + 18 > len(data):
                break
            raw = data[off:off + 15].decode("ascii", "ignore").strip()
            flags = data[off + 15:off + 18]
            # نوع 0x00 = Workstation، والبت الأعلى = group
            group = bool(flags[1] & 0x80) if len(flags) > 1 else False
            if raw and not group and raw not in ("__MSBROWSE__",):
                return raw
            off += 18
    except Exception:
        pass
    return ""


def name_mdns(ip, timeout=0.4):
    """
    استعلام mDNS reverse (UDP 5353) — بيجيب أسماء أجهزة أبل/أندرويد/لينكس
    اللي بتعلن عن نفسها (xxx.local).
    """
    try:
        rev = ".".join(reversed(str(ip).split("."))) + ".in-addr.arpa"
        tid = b"\x00\x00"
        header = tid + b"\x00\x00\x00\x01\x00\x00\x00\x00\x00\x00"
        q = b""
        for part in rev.split("."):
            q += bytes([len(part)]) + part.encode()
        q += b"\x00" + b"\x00\x0c" + b"\x00\x01"  # PTR, IN
        pkt = header + q

        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
        s.settimeout(timeout)
        s.sendto(pkt, ("224.0.0.251", 5353))
        data, _ = s.recvfrom(2048)
        s.close()

        # نبحث عن أول اسم في قسم الإجابات (نقرأ labels)
        name = _parse_first_ptr(data)
        if name:
            return name.replace(".local", "").strip(". ")
    except Exception:
        pass
    return ""


def _parse_first_ptr(data):
    """استخراج أول اسم PTR من رد mDNS (مبسّط، يتعامل مع الضغط)."""
    try:
        # نتخطّى الهيدر والسؤال
        off = 12
        while off < len(data) and data[off] != 0:
            off += data[off] + 1
        off += 1 + 4  # null + qtype + qclass
        # أول إجابة: name(مضغوط غالبًا) + type + class + ttl + rdlen + rdata
        # نتخطّى الاسم
        if data[off] & 0xC0 == 0xC0:
            off += 2
        else:
            while off < len(data) and data[off] != 0:
                off += data[off] + 1
            off += 1
        off += 2 + 2 + 4  # type + class + ttl
        rdlen = struct.unpack(">H", data[off:off + 2])[0]
        off += 2
        # نقرأ الـ rdata كـ labels
        labels = []
        end = off + rdlen
        while off < end and off < len(data):
            length = data[off]
            if length == 0:
                break
            if length & 0xC0 == 0xC0:  # مؤشر ضغط
                ptr = struct.unpack(">H", data[off:off + 2])[0] & 0x3FFF
                off2 = ptr
                while off2 < len(data) and data[off2] != 0:
                    ln = data[off2]
                    if ln & 0xC0 == 0xC0:
                        break
                    labels.append(data[off2 + 1:off2 + 1 + ln].decode("utf-8", "ignore"))
                    off2 += ln + 1
                break
            labels.append(data[off + 1:off + 1 + length].decode("utf-8", "ignore"))
            off += length + 1
        return ".".join(labels)
    except Exception:
        return ""


def resolve_name(ip, do_netbios=True, do_mdns=True):
    """يجرّب المصادر الثلاثة بالترتيب ويرجّع أول اسم."""
    n = name_reverse_dns(ip)
    if n:
        return n
    if do_netbios:
        n = name_netbios(ip)
        if n:
            return n
    if do_mdns:
        n = name_mdns(ip)
        if n:
            return n
    return ""


# ---------------------------------------------------------------------------
#  كشف المشاكل: IP متكرر / MAC مكرر / MAC متشابه / MAC عشوائي
# ---------------------------------------------------------------------------

def is_random_mac(mac):
    """
    الـ MAC العشوائي/المحلي: البت الثاني في أول خانة (locally administered).
    الموبايلات الحديثة بتستخدمه للخصوصية.
    """
    try:
        first = int(mac.split(":")[0], 16)
        return bool(first & 0x02)
    except Exception:
        return False


def analyze_flags(devices):
    """
    يضيف لكل جهاز قائمة flags بالمشاكل المكتشفة.
    يعدّل devices in-place ويرجّعها.
    """
    from collections import defaultdict

    by_mac = defaultdict(list)
    by_ip = defaultdict(list)
    for d in devices:
        if d.get("mac") and d["mac"] != "—":
            by_mac[d["mac"]].append(d["ip"])
        by_ip[d["ip"]].append(d.get("mac", ""))

    # ماكات متشابهة: نفس أول 5 خانات (تختلف في الأخيرة بس)
    prefix5 = defaultdict(set)
    for d in devices:
        mac = d.get("mac", "")
        if mac and mac != "—" and len(mac) == 17:
            prefix5[mac[:14]].add(mac)

    for d in devices:
        flags = []
        mac = d.get("mac", "")
        ip = d["ip"]

        # IP متكرر: نفس الـ IP رد بأكتر من MAC مختلف
        macs_here = set(m for m in by_ip[ip] if m and m != "—")
        if len(macs_here) > 1:
            flags.append("IP متكرر")

        if mac and mac != "—":
            # MAC مكرر: نفس الـ MAC على أكتر من IP
            if len(set(by_mac[mac])) > 1:
                flags.append("MAC مكرر")
            # MAC متشابه: نفس البادئة مع اختلاف الخانة الأخيرة
            if len(prefix5.get(mac[:14], set())) > 1:
                flags.append("MAC متشابه")
            # MAC عشوائي
            if is_random_mac(mac):
                flags.append("MAC عشوائي")

        d["flags"] = flags
    return devices


# ---------------------------------------------------------------------------
#  المسح الكامل
# ---------------------------------------------------------------------------

def scan_network(network, progress_cb=None, result_cb=None, stop_flag=None,
                 do_names=True, do_vendor=True, do_netbios=True, do_mdns=True,
                 workers=None):
    """
    يمسح الشبكة بالكامل ويرجّع list من dict:
      {ip, mac, vendor, name, is_me, flags}
    progress_cb(done, total) و result_cb(dict) اختياريين.
    """
    if stop_flag is None:
        stop_flag = threading.Event()

    hosts = list(network.hosts())
    total = len(hosts)
    if total == 0:
        return []

    # عدد الخيوط الديناميكي حسب حجم الشبكة
    if workers is None:
        workers = min(256, max(50, total // 4))

    oui = load_oui() if do_vendor else {}
    done = [0]
    alive = []
    lock = threading.Lock()
    q = queue.Queue()
    for h in hosts:
        q.put(h)

    def worker():
        while not stop_flag.is_set():
            try:
                ip = q.get_nowait()
            except queue.Empty:
                return
            up = stir(ip)
            with lock:
                done[0] += 1
                if progress_cb:
                    progress_cb(done[0], total)
                if up:
                    alive.append(str(ip))
            q.task_done()

    threads = [threading.Thread(target=worker, daemon=True)
               for _ in range(min(workers, total))]
    for t in threads:
        t.start()
    for t in threads:
        t.join()

    if stop_flag.is_set():
        return []

    time.sleep(1.0)
    arp_pairs = read_arp_table()

    # ندمج الأجهزة الحيّة + جدول ARP
    # نحتفظ بكل أزواج (ip, mac) عشان نكشف التكرار
    all_pairs = []
    seen = set()
    for ip, mac in arp_pairs:
        try:
            if ipaddress.ip_address(ip) in network:
                key = (ip, mac)
                if key not in seen:
                    seen.add(key)
                    all_pairs.append((ip, mac))
        except Exception:
            continue

    # الأجهزة الحيّة اللي ملهاش MAC في ARP (غالبًا على أندرويد)
    ips_with_mac = set(ip for ip, _ in all_pairs)
    for ip in alive:
        if ip not in ips_with_mac:
            all_pairs.append((ip, ""))

    all_pairs.sort(key=lambda x: (int(ipaddress.ip_address(x[0])), x[1]))

    local_ip = get_local_ip()
    devices = []
    dlock = threading.Lock()

    def enrich(pair):
        ip, mac = pair
        d = {
            "ip": ip,
            "mac": mac or "—",
            "vendor": (vendor_of(mac, oui) if (mac and do_vendor) else "") or "—",
            "name": resolve_name(ip, do_netbios, do_mdns) if do_names else "",
            "is_me": ip == local_ip,
            "flags": [],
        }
        with dlock:
            devices.append(d)

    ths = [threading.Thread(target=enrich, args=(p,), daemon=True)
           for p in all_pairs]
    for t in ths:
        t.start()
    for t in ths:
        t.join()

    devices.sort(key=lambda d: (int(ipaddress.ip_address(d["ip"])), d["mac"]))
    analyze_flags(devices)

    # نستدعي result_cb بعد التحليل (عشان الـ flags تكون جاهزة)
    if result_cb:
        for d in devices:
            result_cb(d)

    return devices


# ---------------------------------------------------------------------------
#  التصدير
# ---------------------------------------------------------------------------

def export_csv(devices, path):
    """يصدّر النتائج لملف CSV (مع BOM عشان العربية تظهر صح في Excel)."""
    import csv
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["IP", "MAC", "الشركة", "اسم الجهاز", "ملاحظات"])
        for d in devices:
            w.writerow([d["ip"], d["mac"], d.get("vendor", ""),
                        d.get("name", ""), " / ".join(d.get("flags", []))])
    return path


def export_xlsx(devices, path):
    """
    يصدّر لملف Excel لو openpyxl متوفرة، وإلا يرجع None (نستخدم CSV).
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except Exception:
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = "الأجهزة"
    ws.sheet_view.rightToLeft = True

    headers = ["IP", "MAC", "الشركة", "اسم الجهاز", "ملاحظات"]
    ws.append(headers)
    hfill = PatternFill("solid", fgColor="2563EB")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hfill
        c.alignment = Alignment(horizontal="center")

    warn = PatternFill("solid", fgColor="FEE2E2")
    for d in devices:
        row = [d["ip"], d["mac"], d.get("vendor", ""), d.get("name", ""),
               " / ".join(d.get("flags", []))]
        ws.append(row)
        if d.get("flags"):
            for c in ws[ws.max_row]:
                c.fill = warn

    widths = [16, 20, 28, 24, 26]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    wb.save(path)
    return path


# توافق مع النسخة القديمة
def guess_network(local_ip, prefix=24):
    try:
        return ipaddress.ip_network(f"{local_ip}/{prefix}", strict=False)
    except Exception:
        return ipaddress.ip_network(f"{local_ip}/24", strict=False)
